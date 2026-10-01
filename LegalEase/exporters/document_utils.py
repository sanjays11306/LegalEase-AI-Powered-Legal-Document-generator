"""Document helpers: parse text into blocks, render HTML preview, export TXT/DOCX/PDF.

Document text conventions (used by the generator, editor and exporters):
  * 1st non-empty line = title, 2nd = subtitle (both centred)
  * ALL-CAPS lines (optionally numbered, e.g. "1. SERVICES") = headings; lines starting with "# " are headings too
    (needed for Tamil / Hindi, which have no capital letters); so are short numbered lines with no Latin lowercase
  * "1.1 text" and "A. text" lines = clauses / recitals (hanging indent)
  * lines starting with "- " = bullets
  * **bold** and *italic* inline markup
  * [Bracketed text] and ____ blanks = placeholders (highlighted so they are not forgotten);
    blanks on Signature / Name / Date lines are meant for pen and are not flagged.
"""
import html
import io
import re
from collections import Counter
from functools import lru_cache
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
PLACEHOLDER = re.compile(r"\[[^\]\n]{1,60}\]|_{4,}")
# Signature / name / date lines meant for pen: "Signature: ____", "1. Name: ____  Signature: ____" in ANY language
SIG_LINE = re.compile(r"^(\d+\.\s*)?(Signature|Name|Date)\s*:|^(\d+\.\s*)?([^:_\n]{1,40}:\s*_{3,}\s*)+$", re.I)
INDIC = re.compile("[\u0900-\u0DFF]")   # Devanagari, Bengali, Gurmukhi, Gujarati, Odia, Tamil, Telugu, Kannada, Malayalam
CLAUSE = re.compile(r"^(\d+\.\d+)\s+(.*)$")
RECITAL = re.compile(r"^([A-Z]\.)\s+(.*)$")
KEEP_TOGETHER = re.compile(r"^(For and on behalf|Signature:|Name:|IN WITNESS)", re.I)


# ------------------------------------------------------------------ parsing
def parse_document(text: str) -> list[tuple[str, str]]:
    """Turn raw text into a list of (kind, text) blocks."""
    blocks, seen = [], 0
    for raw in text.replace("\r", "").split("\n"):
        line = raw.strip()
        if not line:
            blocks.append(("gap", ""))
            continue
        seen += 1
        explicit = line.startswith("# ")
        if explicit:
            line = line[2:].strip()
        core = re.sub(r"^\d+\.\s*", "", line).replace("**", "")
        if seen == 1:
            blocks.append(("title", line))
        elif seen == 2:
            blocks.append(("subtitle", line))
        elif explicit:
            blocks.append(("heading", line))
        elif CLAUSE.match(line) or (RECITAL.match(line) and not RECITAL.match(line).group(2).replace("**", "").isupper()):
            blocks.append(("clause", line))
        elif line.startswith("- "):
            blocks.append(("bullet", line[2:]))
        elif core.isupper() and len(core) < 70 and any(c.isalpha() for c in core):
            blocks.append(("heading", line))
        elif (re.match(r"^\d+\.\s+\S", line) and len(core) < 70 and not re.search(r"[a-z:_]|[.;,\u0964]$", core)
              and INDIC.search(core)):     # "3. <Tamil/Hindi heading>": uncased scripts cannot use isupper()
            blocks.append(("heading", line))
        else:
            blocks.append(("para", line))
    return blocks


def split_label(t: str) -> tuple[str, str]:
    """'1.1 Text' -> ('1.1', 'Text');  'A. Text' -> ('A.', 'Text');  otherwise ('', t)."""
    m = CLAUSE.match(t) or RECITAL.match(t)
    return (m.group(1), m.group(2)) if m else ("", t)


def find_placeholders(text: str) -> list[str]:
    """All unfilled placeholders (signature-line blanks are intentional and ignored)."""
    found = []
    for line in text.replace("\r", "").split("\n"):
        if not SIG_LINE.match(line.strip()):
            found += PLACEHOLDER.findall(line)
    return found


def placeholder_report(text: str) -> list[tuple[str, int]]:
    """[(label, count)] for the 'items still need your input' banner."""
    counts = Counter(t if t.startswith("[") else "blank line (____)" for t in find_placeholders(text))
    return sorted(counts.items(), key=lambda kv: kv[0])


def suggest_filename(text: str, ext: str) -> str:
    """'EMPLOYMENT CONTRACT' -> 'employment_contract.<ext>' (falls back to legalease_document)."""
    first = next((l.strip() for l in text.split("\n") if l.strip()), "")
    slug = re.sub(r"[^a-z0-9]+", "_", first.replace("*", "").lower()).strip("_")[:50] or "legalease_document"
    return f"{slug}.{ext}"


# ------------------------------------------------------------------ HTML preview
def _inline_html(s: str, mark: bool = False) -> str:
    s = html.escape(s)
    if mark:
        s = PLACEHOLDER.sub(lambda m: f'<mark class="ph">{m.group(0)}</mark>', s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    return re.sub(r"\*(.+?)\*", r"<i>\1</i>", s)


def render_html(text: str, short: bool = False) -> str:
    """Return the scrollable 'paper' preview as an HTML string (no blank lines)."""
    css = {"title": "t", "subtitle": "s", "heading": "h", "para": "p", "bullet": "b", "clause": "c"}
    parts = []
    for kind, t in parse_document(text):
        if kind == "gap":
            continue
        mark = not SIG_LINE.match(t)
        if kind == "clause":
            label, body = split_label(t)
            inner = (f'<span class="n">{html.escape(label)}</span>' if label else "") + _inline_html(body, mark)
        else:
            inner = ("&bull; " if kind == "bullet" else "") + _inline_html(t, mark)
        parts.append(f'<div class="{css[kind]}">{inner}</div>')
    cls = "le-paper short" if short else "le-paper"
    return f'<div class="le-scroll"><div class="{cls}">{"".join(parts)}</div></div>'


# ------------------------------------------------------------------ TXT
def to_txt(text: str) -> bytes:
    """Plain-text export (markup removed)."""
    plain = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    plain = re.sub(r"\*(.+?)\*", r"\1", plain)
    plain = re.sub(r"^- ", "\u2022 ", plain, flags=re.M)
    return plain.replace("\r\n", "\n").encode("utf-8")


# ------------------------------------------------------------------ DOCX
def _add_field(par, code: str, size: Pt) -> None:
    run = par.add_run()
    run.font.size = size
    begin, instr, end = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def _bottom_border(par) -> None:
    border, bottom = OxmlElement("w:pBdr"), OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "8"), ("space", "6"), ("color", "222222")):
        bottom.set(qn(f"w:{k}"), v)
    border.append(bottom)
    par._p.get_or_add_pPr().append(border)


_RPR_AFTER_BCS = ("w:i", "w:iCs", "w:caps", "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow", "w:emboss",
                  "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish", "w:webHidden", "w:color", "w:spacing", "w:w",
                  "w:kern", "w:position", "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
                  "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang", "w:eastAsianLayout", "w:specVanish")


def _complex_run(run, bold: bool, size_pt: int | None) -> None:
    """Mark a Tamil/Hindi/... run as complex-script so Word uses the complex-script font and bold/size settings."""
    run.font.complex_script = True
    rpr = run._r.get_or_add_rPr()
    if bold:
        rpr.insert_element_before(OxmlElement("w:bCs"), *_RPR_AFTER_BCS)
    if size_pt:
        sz = OxmlElement("w:szCs")
        sz.set(qn("w:val"), str(size_pt * 2))
        rpr.insert_element_before(sz, *_RPR_AFTER_BCS[_RPR_AFTER_BCS.index("w:szCs") + 1:])


def _add_runs(p, t: str, kind: str, mark: bool) -> None:
    always_bold = kind in ("title", "subtitle", "heading")
    for piece in INLINE.split(t):
        if not piece:
            continue
        bold = piece.startswith("**") and piece.endswith("**") and len(piece) > 4
        ital = not bold and piece.startswith("*") and piece.endswith("*") and len(piece) > 2
        txt = piece[2:-2] if bold else piece[1:-1] if ital else piece
        for part in re.split(f"({PLACEHOLDER.pattern})", txt):
            if not part:
                continue
            run = p.add_run(part)
            run.bold, run.italic = bold or always_bold, ital
            if kind == "title":
                run.font.size = Pt(16)
            if INDIC.search(part):
                _complex_run(run, bold or always_bold, 16 if kind == "title" else None)
            if mark and PLACEHOLDER.fullmatch(part):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW  # easy to spot in Word


def to_docx(text: str) -> bytes:
    """Word export: hanging-indent clauses, real bullets, kept-together headings/signatures,
    highlighted placeholders and 'Page X of Y' footer."""
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    base = doc.styles["Normal"]
    base.font.name, base.font.size = "Times New Roman", Pt(12)
    base.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    base.element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Nirmala UI")   # Tamil / Hindi / other Indian scripts (ships with Windows)
    base.paragraph_format.space_after = Pt(6)
    base.paragraph_format.line_spacing = 1.15
    doc.core_properties.title = next((l.strip().strip("*") for l in text.split("\n") if l.strip()), "Legal Document")
    doc.core_properties.author = "LegalEase"

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run("Page ").font.size = Pt(9)
    _add_field(foot, "PAGE", Pt(9))
    foot.add_run(" of ").font.size = Pt(9)
    _add_field(foot, "NUMPAGES", Pt(9))

    for kind, t in parse_document(text):
        if kind == "gap":
            continue
        mark = not SIG_LINE.match(t)
        p = doc.add_paragraph(style="List Bullet") if kind == "bullet" else doc.add_paragraph()
        fmt = p.paragraph_format
        if kind in ("title", "subtitle"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if kind == "subtitle":
                fmt.space_after = Pt(14)
                _bottom_border(p)
        elif kind == "para":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "clause":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            label, body = split_label(t)
            if label:
                fmt.left_indent, fmt.first_line_indent = Inches(0.5), Inches(-0.5)
                fmt.tab_stops.add_tab_stop(Inches(0.5))
                t = f"{label}\t{body}"
        elif kind == "heading":
            fmt.space_before, fmt.keep_with_next = Pt(14), True
        if KEEP_TOGETHER.match(t):
            fmt.keep_with_next = True
        _add_runs(p, t, kind, mark)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
_FONT_SETS = [  # (regular, bold, italic, bold-italic): first complete set found on this machine wins
    ("C:/Windows/Fonts/times.ttf", "C:/Windows/Fonts/timesbd.ttf", "C:/Windows/Fonts/timesi.ttf", "C:/Windows/Fonts/timesbi.ttf"),
    tuple(f"/System/Library/Fonts/Supplemental/Times New Roman{s}.ttf" for s in ("", " Bold", " Italic", " Bold Italic")),
    tuple(f"/usr/share/fonts/truetype/liberation/LiberationSerif-{s}.ttf" for s in ("Regular", "Bold", "Italic", "BoldItalic")),
    tuple(f"/usr/share/fonts/truetype/dejavu/DejaVuSerif{s}.ttf" for s in ("", "-Bold", "-Italic", "-BoldItalic")),
]
_BUILTIN = ("Times-Roman", "Times-Bold", "Times-Italic", "Times-BoldItalic")


@lru_cache(maxsize=1)
def _pdf_fonts() -> tuple[tuple[str, str, str, str], bool]:
    """Register a Unicode serif font if available. Returns (font names, supports the rupee sign)."""
    for paths in _FONT_SETS:
        if not all(Path(p).exists() for p in paths):
            continue
        try:
            names = ("LESerif", "LESerif-B", "LESerif-I", "LESerif-BI")
            for n, p in zip(names, paths):
                pdfmetrics.registerFont(TTFont(n, p))
            pdfmetrics.registerFontFamily("LESerif", normal=names[0], bold=names[1], italic=names[2], boldItalic=names[3])
            return names, 0x20B9 in pdfmetrics.getFont(names[0]).face.charToGlyph
        except Exception:
            continue
    return _BUILTIN, False


class _NumberedCanvas(rl_canvas.Canvas):
    """Canvas that writes 'Page X of Y' (needs the total, so pages are saved first)."""

    def __init__(self, *a, **k):
        super().__init__(*a, **k)
        self._saved = []

    def showPage(self):
        self._saved.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved)
        for state in self._saved:
            self.__dict__.update(state)
            self.setFont("Times-Roman", 9)
            self.drawCentredString(A4[0] / 2, 0.5 * inch, f"Page {self._pageNumber} of {total}")
            super().showPage()
        super().save()


@lru_cache(maxsize=1)
def _weasyprint_ok() -> bool:
    """WeasyPrint (Pango + HarfBuzz) shapes Indian scripts correctly. ReportLab does NOT: vowel signs and
    conjuncts come out in the wrong place even with a perfect font."""
    try:
        import weasyprint
        weasyprint.HTML(string="<p>x</p>").write_pdf()
        return True
    except Exception:
        return False


def pdf_support(text: str) -> tuple[bool, str]:
    """(can we make a correct PDF of this text?, reason if not)."""
    if not INDIC.search(text):
        return True, ""
    if _weasyprint_ok():
        return True, ""
    return False, ("PDF export of Tamil / Hindi needs WeasyPrint and a Noto font installed on the server "
                   "(see README). Download the DOCX instead - Word shapes these scripts correctly.")


_PDF_CSS = """@page { size: A4; margin: 25.4mm; @bottom-center { content: "Page " counter(page) " of " counter(pages); font-size: 9pt; } }
body { font-family: "Noto Serif Tamil", "Noto Serif Devanagari", "Noto Serif", "Noto Sans Tamil", "Noto Sans Devanagari",
       "Nirmala UI", "FreeSerif", serif; font-size: 11.5pt; line-height: 1.5; text-align: justify; }
.t { text-align: center; font-weight: 700; font-size: 17pt; margin: 0 0 4pt; }
.s { text-align: center; font-weight: 700; margin: 0 0 12pt; padding-bottom: 6pt; border-bottom: 1px solid #222; }
.h { font-weight: 700; margin: 12pt 0 4pt; text-align: left; break-after: avoid; }
.p, .b, .c { margin: 0 0 5pt; }
.c, .b { padding-left: 34pt; text-indent: -34pt; } .n { display: inline-block; width: 34pt; text-indent: 0; }"""


def _to_pdf_weasyprint(text: str) -> bytes:
    from weasyprint import HTML
    css = {"title": "t", "subtitle": "s", "heading": "h", "para": "p", "bullet": "b", "clause": "c"}
    parts = []
    for kind, t in parse_document(text):
        if kind == "gap":
            continue
        if kind == "clause":
            label, body = split_label(t)
            inner = (f'<span class="n">{html.escape(label)}</span>' if label else "") + _inline_html(body)
        else:
            inner = ("&bull;&nbsp;" if kind == "bullet" else "") + _inline_html(t)
        parts.append(f'<div class="{css[kind]}">{inner}</div>')
    doc = f'<!doctype html><html><head><meta charset="utf-8"><style>{_PDF_CSS}</style></head><body>{"".join(parts)}</body></html>'
    return HTML(string=doc).write_pdf()


def to_pdf(text: str) -> bytes:
    """PDF export. Latin text: ReportLab (Unicode font when available, else built-in Times and 'Rs.').
    Tamil / Hindi / other Indian scripts: WeasyPrint, because ReportLab cannot shape them."""
    if INDIC.search(text):
        ok, why = pdf_support(text)
        if not ok:
            raise RuntimeError(why)
        return _to_pdf_weasyprint(text)
    (reg, bold, _, _), rupee = _pdf_fonts()
    if not rupee:
        text = text.replace("\u20b9", "Rs. ")
    base = ParagraphStyle("b", fontName=reg, fontSize=11.5, leading=16, spaceAfter=5, alignment=TA_JUSTIFY,
                          bulletFontName=reg, bulletFontSize=11.5)
    styles = {
        "title": ParagraphStyle("t", parent=base, fontName=bold, fontSize=17, alignment=TA_CENTER, leading=22, spaceAfter=4),
        "subtitle": ParagraphStyle("s", parent=base, fontName=bold, alignment=TA_CENTER, spaceAfter=6),
        "heading": ParagraphStyle("h", parent=base, fontName=bold, spaceBefore=12, alignment=TA_LEFT, keepWithNext=1),
        "para": base,
        "clause": ParagraphStyle("c", parent=base, leftIndent=34, bulletIndent=0),
        "bullet": ParagraphStyle("bu", parent=base, leftIndent=34, bulletIndent=16),
    }
    story = []
    for kind, t in parse_document(text):
        if kind == "gap":
            continue
        label = ""
        if kind == "clause":
            label, t = split_label(t)
        markup = _inline_html(t)
        if kind in ("title", "subtitle", "heading"):
            markup = f"<b>{markup}</b>"
        style = styles[kind]
        if kind in ("para", "clause", "bullet") and KEEP_TOGETHER.match(t):
            style = ParagraphStyle("k", parent=style, keepWithNext=1)
        story.append(Paragraph(markup, style, bulletText=("\u2022" if kind == "bullet" else label) or None))
        if kind == "subtitle":
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#222222"), spaceAfter=10))
    if not story:
        story.append(Spacer(1, 0.1 * inch))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch,
                      title="LegalEase Document", author="LegalEase").build(story, canvasmaker=_NumberedCanvas)
    return buf.getvalue()
